# -*- coding: utf-8 -*-
"""Excel 填表助手

功能概览
--------
* 记住多个常用的表格路径，随时切换；列表上右键可以移除记录（不会删除文件本身）；
* 一键用系统默认程序打开表格，但填写时**不需要**把表格开到前台；
* 指定工作表与行号（支持 5 / 5,7 / 5-8 这类写法）；
* 切到某一行时，如果该行某些列本来就有内容，输入框里会直接显示出来；
  只有你改动过的列才会被写回表格，没碰过的单元格（公式、格式）保持原样；
* 默认 3 个输入框，可自由增减；输入框从左到右依次对应从左到右的列；
  输入框太多时自动换行，窗口跟着放大缩小；
* 「← / →」整体平移这一排输入框所对应的列，平移前填过的内容不会丢，只是暂时不显示；
  平移到更高的列位可以继续填；
* 输入框获得焦点时，里面的内容会自动全选，直接输入即可覆盖；
* 「写入表格」一键写入（不弹二次确认），在任意输入框里按回车也能写入；
* 「取消」放弃本次改动，把输入框恢复成表格里原来的内容；
* 切换行/表/工作表时若还有没写入的改动，会问你要不要先写入；
  也可以在右下角把「切换行时」设成“自动写入”或“自动放弃”。

支持的文件：.xlsx / .xlsm / .xltx / .xltm（openpyxl 直接读写）
            .xls / .xlsb / .et（需要本机装有 Excel 或 WPS，走 COM 接口，能保留原始格式）
"""

from __future__ import annotations

import datetime as _dt
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except Exception:  # openpyxl 缺失时界面仍可用，只是不能读写 xlsx
    openpyxl = None

    def get_column_letter(n: int) -> str:
        s = ""
        while n > 0:
            n, r = divmod(n - 1, 26)
            s = chr(65 + r) + s
        return s


# ---------------------------------------------------------------- 常量

APP_TITLE = "Excel 填表助手"
if getattr(sys, "frozen", False):          # 打包成 exe 后，配置放在 exe 旁边
    APP_DIR = Path(sys.executable).resolve().parent
else:
    APP_DIR = Path(__file__).resolve().parent
CONFIG_FILE = APP_DIR / "config.json"
FALLBACK_CONFIG = Path.home() / ".excel_filler_config.json"

DEFAULT_COUNT = 3      # 默认输入框数量
MAX_COUNT = 24         # 输入框数量上限
PER_ROW = 6            # 每行最多放几个输入框，超了自动换行
MAX_COL = 16384        # Excel 最大列数（XFD）
READ_MAX_COL = 500     # 读取已有内容时最多读多少列
MIN_W, MIN_H = 700, 540

XLSX_EXTS = {".xlsx", ".xlsm", ".xltx", ".xltm"}
COM_EXTS = {".xls", ".xlsb", ".xlt", ".et", ".ett"}

SWITCH_LABELS = ["询问", "自动写入", "自动放弃"]
SWITCH_KEYS = {"询问": "ask", "自动写入": "write", "自动放弃": "discard"}
SWITCH_LABEL_OF = {v: k for k, v in SWITCH_KEYS.items()}

FONT = ("Microsoft YaHei UI", 10)
FONT_BOLD = ("Microsoft YaHei UI", 10, "bold")
FONT_SMALL = ("Microsoft YaHei UI", 9)
FONT_TITLE = ("Microsoft YaHei UI", 15, "bold")
FONT_BADGE = ("Microsoft YaHei UI", 9, "bold")

BG = "#f4f5f7"
FG = "#1f2937"
MUTED = "#6b7280"
ACCENT = "#2563eb"


# ---------------------------------------------------------------- 工具函数

def parse_rows(text: str) -> list[int]:
    """把 '5' / '5,7' / '5-8' 这类写法解析成行号列表。"""
    rows: set[int] = set()
    text = (text or "").strip()
    if not text:
        return []
    for part in re.split(r"[,，;；\s]+", text):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)\s*[-~～－—]\s*(\d+)", part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a > b:
                a, b = b, a
            rows.update(range(a, b + 1))
        elif part.isdigit():
            rows.add(int(part))
        else:
            raise ValueError(part)
    return sorted(r for r in rows if r >= 1)


def rows_text(rows: list[int]) -> str:
    """把行号列表压成好看的字符串，如 5,7-9。"""
    if not rows:
        return ""
    parts, start, prev = [], rows[0], rows[0]
    for r in rows[1:]:
        if r == prev + 1:
            prev = r
            continue
        parts.append(f"{start}-{prev}" if start != prev else f"{start}")
        start = prev = r
    parts.append(f"{start}-{prev}" if start != prev else f"{start}")
    return ",".join(parts)


def format_cell(v) -> str:
    """把单元格的值转成输入框里能显示的文字。"""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "TRUE" if v else "FALSE"
    if isinstance(v, _dt.datetime):
        if v.hour or v.minute or v.second:
            return v.strftime("%Y-%m-%d %H:%M:%S")
        return v.strftime("%Y-%m-%d")
    if isinstance(v, _dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return str(int(v)) if v.is_integer() else str(v)
    return str(v).strip()


def coerce_value(text: str):
    """数字自动转成数值，其余保持文本；前导零（007）和超长数字保持文本。"""
    t = text.strip()
    if not t:
        return None
    if re.fullmatch(r"[+-]?0\d+", t):
        return text
    if re.fullmatch(r"[+-]?\d+", t):
        if len(t.lstrip("+-")) >= 15:      # 超出 Excel 的 15 位精度，保持文本
            return text
        try:
            return int(t)
        except ValueError:
            return text
    if re.fullmatch(r"[+-]?(?:\d+\.\d*|\.\d+)(?:[eE][+-]?\d+)?", t) or \
       re.fullmatch(r"[+-]?\d+[eE][+-]?\d+", t):
        try:
            return float(t)
        except ValueError:
            return text
    return text


# ---------------------------------------------------------------- Excel / WPS 通信

PS_COMMON = r"""
$ErrorActionPreference = 'Stop'
$raw  = [System.IO.File]::ReadAllText($JsonPath, [System.Text.Encoding]::UTF8)
$data = $raw | ConvertFrom-Json

$excel = $null
foreach ($progid in @('Excel.Application', 'Ket.Application', 'KET.Application')) {
    try { $excel = New-Object -ComObject $progid; break } catch { $excel = $null }
}
if ($null -eq $excel) { Write-Error 'NO_EXCEL'; exit 3 }
$excel.Visible = $false
$excel.DisplayAlerts = $false
"""

PS_WRITE = r"""
param([Parameter(Mandatory=$true)][string]$JsonPath)
""" + PS_COMMON + r"""
try {
    $wb = $excel.Workbooks.Open($data.path)
    $ws = $wb.Worksheets.Item($data.sheet)
    foreach ($c in $data.cells) {
        if ($null -eq $c.value) { $ws.Cells.Item([int]$c.row, [int]$c.col).ClearContents() }
        else { $ws.Cells.Item([int]$c.row, [int]$c.col).Value2 = $c.value }
    }
    $wb.Save()
    $wb.Close($false)
    Write-Output 'OK'
} catch {
    Write-Error $_.Exception.Message
    exit 4
} finally {
    try { $excel.Quit() } catch {}
}
"""

PS_READ = r"""
param([Parameter(Mandatory=$true)][string]$JsonPath)
""" + PS_COMMON + r"""
try {
    $wb = $excel.Workbooks.Open($data.path, 0, $true)
    $ws = $wb.Worksheets.Item($data.sheet)
    $last = $ws.UsedRange.Column + $ws.UsedRange.Columns.Count - 1
    if ($last -gt $data.maxCol) { $last = $data.maxCol }
    $result = @{}
    for ($c = 1; $c -le $last; $c++) {
        $t = [string]$ws.Cells.Item($data.row, $c).Text
        if ($t.Trim() -ne '') { $result[[string]$c] = $t.Trim() }
    }
    $wb.Close($false)
    [System.IO.File]::WriteAllText($data.out, ($result | ConvertTo-Json -Compress), (New-Object System.Text.UTF8Encoding($false)))
    Write-Output 'OK'
} catch {
    Write-Error $_.Exception.Message
    exit 4
} finally {
    try { $excel.Quit() } catch {}
}
"""


def _run_ps(script: str, payload: dict) -> tuple[bool, str, Path | None]:
    """把 payload 交给一段 PowerShell 脚本执行，返回 (是否成功, 信息, 临时目录)。"""
    powershell = shutil.which("powershell") or shutil.which("pwsh")
    if not powershell:
        return False, "未找到 PowerShell，无法调用 Excel/WPS。", None

    tmpdir = Path(tempfile.mkdtemp(prefix="fill_excel_"))
    js = tmpdir / "payload.json"
    ps = tmpdir / "task.ps1"
    js.write_text(json.dumps(payload, ensure_ascii=False), encoding="utf-8")
    ps.write_text(script, encoding="utf-8")
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    try:
        p = subprocess.run(
            [powershell, "-NoProfile", "-ExecutionPolicy", "Bypass",
             "-File", str(ps), "-JsonPath", str(js)],
            stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
            encoding="utf-8", errors="replace", timeout=300, creationflags=flags)
    except subprocess.TimeoutExpired:
        shutil.rmtree(tmpdir, ignore_errors=True)
        return False, "调用 Excel 超时（超过 5 分钟）。", None

    if p.returncode == 0 and "OK" in (p.stdout or ""):
        return True, "", tmpdir
    err = (p.stderr or "") + (p.stdout or "")
    shutil.rmtree(tmpdir, ignore_errors=True)
    if "NO_EXCEL" in err:
        return False, ("本机没有检测到 Excel 或 WPS，无法直接处理该格式。\n"
                       "请用 Excel/WPS 打开后「另存为」.xlsx，再重新添加该文件。"), None
    return False, "操作失败：" + err.strip()[:400], None


def com_write(path: str, sheet: str, cells: list[dict]) -> tuple[bool, str]:
    ok, info, _ = _run_ps(PS_WRITE, {"path": str(Path(path).resolve()), "sheet": sheet, "cells": cells})
    return (True, "已通过 Excel/WPS 写入。") if ok else (False, info)


def com_read(path: str, sheet: str, row: int) -> tuple[dict[int, str] | None, str]:
    tmp_out = Path(tempfile.gettempdir()) / f"fill_excel_read_{os.getpid()}.json"
    ok, info, tmpdir = _run_ps(PS_READ, {"path": str(Path(path).resolve()), "sheet": sheet,
                                         "row": row, "maxCol": READ_MAX_COL, "out": str(tmp_out)})
    if not ok or tmpdir is None:
        return None, info
    try:
        raw = json.loads(tmp_out.read_text(encoding="utf-8") or "{}")
        data = {int(k) - 1: str(v).strip() for k, v in raw.items() if str(v).strip()}
        return data, ""
    except Exception as e:
        return None, f"读取结果解析失败：{e}"
    finally:
        shutil.rmtree(tmpdir, ignore_errors=True)
        tmp_out.unlink(missing_ok=True)


def read_row_xlsx(path: str, sheet: str, row: int) -> tuple[dict[int, str] | None, str]:
    """读取 xlsx 某一行已有的内容，返回 {0 基列号: 文本}。"""
    if openpyxl is None:
        return None, "缺少 openpyxl，无法读取 .xlsx。"
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        return None, f"打开文件失败：{e}"
    try:
        if sheet and sheet in wb.sheetnames:
            ws = wb[sheet]
        else:
            ws = wb.active
        limit = min(max(ws.max_column or 1, 1), READ_MAX_COL)
        data: dict[int, str] = {}
        for r in ws.iter_rows(min_row=row, max_row=row, max_col=limit, values_only=True):
            for i, v in enumerate(r):
                s = format_cell(v)
                if s:
                    data[i] = s
        return data, ""
    except Exception as e:
        return None, f"读取失败：{e}"
    finally:
        try:
            wb.close()
        except Exception:
            pass


def load_config() -> dict:
    for p in (CONFIG_FILE, FALLBACK_CONFIG):
        try:
            data = json.loads(p.read_text(encoding="utf-8"))
        except Exception:
            continue
        return {
            "files": [str(f) for f in data.get("files", []) if isinstance(f, str)],
            "last": str(data.get("last") or ""),
            "switch": str(data.get("switch") or "ask"),
            "window": str(data.get("window") or ""),
        }
    return {"files": [], "last": "", "switch": "ask", "window": ""}


# ---------------------------------------------------------------- 主界面

class FillApp:
    def __init__(self, root: tk.Tk):
        self.root = root
        cfg = load_config()
        self.files: list[str] = cfg["files"]
        self.last: str = cfg["last"]
        self._switch_key: str = cfg.get("switch", "ask")
        self.saved_geometry: str = cfg.get("window", "")   # 上次关窗口时的大小/位置

        self.values: dict[int, str] = {}    # 输入框里的内容（含未显示的列）
        self.base: dict[int, str] = {}      # 表格里该行本来的内容
        self.offset = 0                     # 最左边输入框对应的列（0 基）
        self.count = DEFAULT_COUNT          # 当前输入框数量
        self.vars: list[tuple[int, tk.StringVar]] = []
        self.entries: list[ttk.Entry] = []   # 和 vars 一一对应，用于左右键跳转
        self._focus_no_select = None         # 用键盘跳过去时，别触发“自动全选”
        self._loading = False               # 重建输入框时抑制回调
        self.loaded_path: str | None = None
        self.current_row: int | None = None
        self.row_limit = 500                # 行号下拉可选到多少行
        self.ref_window: RefWindow | None = None
        self.ref_geom: str | None = None    # 参考窗口上次关闭时的位置（只在本次运行有效）
        self._switching = False             # 正在处理“切换行”询问，防止重复弹窗

        self._init_style()
        self._build_ui()
        self._refresh_list()
        self.rebuild_cells()
        self._status("就绪")

    # ---------------------------------------------------------- 样式

    def _init_style(self):
        st = ttk.Style(self.root)
        for theme in ("vista", "winnative", "clam"):
            try:
                st.theme_use(theme)
                break
            except tk.TclError:
                continue
        st.configure(".", font=FONT, background=BG)
        st.configure("TFrame", background=BG)
        st.configure("TLabel", background=BG, foreground=FG, font=FONT)
        st.configure("TLabelframe", background=BG, bordercolor="#d6d9de", relief="solid", borderwidth=1)
        st.configure("TLabelframe.Label", background=BG, foreground=ACCENT, font=FONT_BOLD)
        st.configure("Hint.TLabel", background=BG, foreground=MUTED, font=FONT_SMALL)
        st.configure("Path.TLabel", background=BG, foreground=MUTED, font=FONT_SMALL)
        st.configure("Col.TLabel", background=BG, foreground=ACCENT,
                     font=("Microsoft YaHei UI", 9, "bold"))
        st.configure("Status.TLabel", background=BG, foreground=MUTED, font=FONT_SMALL)
        st.configure("Ok.TLabel", background=BG, foreground="#15803d", font=FONT_SMALL)
        st.configure("Warn.TLabel", background=BG, foreground="#b91c1c", font=FONT_SMALL)
        st.configure("Empty.TLabel", background=BG, foreground="#9ca3af", font=FONT)
        st.configure("TButton", padding=(9, 4))
        st.configure("TEntry", padding=3)
        st.configure("Treeview", background="white", fieldbackground="white",
                     foreground=FG, font=FONT, rowheight=25, borderwidth=0)
        st.configure("Treeview.Heading", font=FONT_SMALL, foreground=MUTED, padding=(6, 4))
        st.map("Treeview", background=[("selected", ACCENT)], foreground=[("selected", "white")])
        st.configure("Ref.Treeview", background="white", fieldbackground="white",
                     foreground=FG, font=FONT, rowheight=24, borderwidth=0)

    # ---------------------------------------------------------- 界面

    def _badge(self, parent, text: str) -> tk.Label:
        """区块标题用的小徽章。"""
        return tk.Label(parent, text=text, font=FONT_BADGE, bg=ACCENT, fg="white",
                        padx=10, pady=2, bd=0)

    def _build_ui(self):
        root = self.root
        root.title(APP_TITLE)
        root.configure(bg=BG)

        main = ttk.Frame(root, padding=(14, 12, 14, 12))
        main.pack(fill="both", expand=True)

        # ---- 顶部标题 ----
        head = ttk.Frame(main)
        head.pack(fill="x", pady=(0, 10))
        tk.Label(head, text=APP_TITLE, font=FONT_TITLE, bg=BG, fg="#111827").pack(side="left")
        tk.Label(head, text="选表格  →  定行  →  填内容  →  写入",
                 font=FONT_SMALL, bg=BG, fg=MUTED).pack(side="left", padx=(16, 0), pady=(7, 0))

        # ---- 1. 表格 ----
        lf_file = ttk.LabelFrame(main, labelwidget=self._badge(main, "1  选择表格"),
                                 padding=(12, 8, 12, 10))
        lf_file.pack(fill="x")

        list_row = ttk.Frame(lf_file)
        list_row.pack(fill="x")
        self.tree_files = ttk.Treeview(list_row, columns=("name", "dir"), show="headings",
                                       height=3, selectmode="browse")
        self.tree_files.heading("name", text="表格")
        self.tree_files.heading("dir", text="所在文件夹")
        self.tree_files.column("name", width=200, anchor="w", stretch=False)
        self.tree_files.column("dir", width=360, anchor="w", stretch=True)
        sb = ttk.Scrollbar(list_row, orient="vertical", command=self.tree_files.yview)
        self.tree_files.configure(yscrollcommand=sb.set)
        self.tree_files.pack(side="left", fill="both", expand=True)
        sb.pack(side="left", fill="y")

        btns = ttk.Frame(list_row)
        btns.pack(side="left", fill="y", padx=(10, 0))
        ttk.Button(btns, text="＋ 添加表格", command=self.add_file, width=12).pack(fill="x")
        ttk.Button(btns, text="打开表格", command=self.open_file, width=12).pack(fill="x", pady=(5, 0))
        ttk.Button(btns, text="✕ 移出列表", command=self.remove_selected, width=12)\
            .pack(fill="x", pady=(5, 0))

        ttk.Label(lf_file, text="提示：在列表上点右键可以打开、定位或移除记录（不会删除文件本身）。",
                  style="Hint.TLabel", anchor="w").pack(fill="x", pady=(6, 0))

        self.tree_files.bind("<<TreeviewSelect>>", self.on_select_file)
        self.tree_files.bind("<Button-3>", self.on_right_click)
        self.tree_files.bind("<Delete>", lambda e: self.remove_selected())

        # ---- 2. 位置 ----
        lf_pos = ttk.LabelFrame(main, labelwidget=self._badge(main, "2  填写位置"),
                                padding=(12, 8, 12, 10))
        lf_pos.pack(fill="x", pady=(10, 0))

        row = ttk.Frame(lf_pos)
        row.pack(fill="x")
        ttk.Label(row, text="工作表：").pack(side="left")
        self.sheet_var = tk.StringVar(value="")
        self.sheet_combo = ttk.Combobox(row, textvariable=self.sheet_var, state="readonly", width=22)
        self.sheet_combo.pack(side="left")
        self.sheet_combo.bind("<<ComboboxSelected>>", lambda e: self.on_sheet_change())
        ttk.Label(row, text="　　行号：").pack(side="left")
        self.row_var = tk.StringVar(value="1")
        self.row_combo = ttk.Combobox(row, textvariable=self.row_var, width=12,
                                      height=12, justify="center",
                                      values=self._row_options())
        self.row_combo.pack(side="left")
        self.row_combo.configure(postcommand=self._reset_row_options)
        self.row_combo.bind("<Return>", self.on_row_commit)
        self.row_combo.bind("<<ComboboxSelected>>", self.on_row_commit)
        self.row_combo.bind("<FocusOut>", self.on_row_commit)
        ttk.Label(row, text="   （点箭头选，也可手填；5,7 / 5-8 写多行）",
                  style="Hint.TLabel").pack(side="left")
        self.auto_num_var = tk.BooleanVar(value=True)
        ttk.Checkbutton(row, text="数字转成数值", variable=self.auto_num_var)\
            .pack(side="right", padx=(12, 0))

        # ---- 3. 列内容 ----
        lf_cell = ttk.LabelFrame(main, labelwidget=self._badge(main, "3  填写内容"),
                                 padding=(12, 8, 12, 10))
        lf_cell.pack(fill="both", expand=True, pady=(10, 0))

        bar = ttk.Frame(lf_cell)
        bar.pack(fill="x")
        self.btn_left = ttk.Button(bar, text="← 左移一列", command=lambda: self.shift(-1), width=11)
        self.btn_left.pack(side="left")
        self.btn_right = ttk.Button(bar, text="右移一列 →", command=lambda: self.shift(1), width=11)
        self.btn_right.pack(side="left", padx=(5, 0))
        self.btn_del = ttk.Button(bar, text="－ 减少", command=self.del_cell, width=9)
        self.btn_del.pack(side="left", padx=(14, 0))
        self.btn_add = ttk.Button(bar, text="＋ 增加", command=self.add_cell, width=9)
        self.btn_add.pack(side="left", padx=(5, 0))
        self.btn_reset = ttk.Button(bar, text="⟲ 复位", command=self.reset_offset, width=9)
        self.btn_reset.pack(side="left", padx=(14, 0))
        self.range_var = tk.StringVar(value="")
        ttk.Label(bar, textvariable=self.range_var, style="Hint.TLabel").pack(side="left", padx=(16, 0))

        self.cells_frame = ttk.Frame(lf_cell)
        self.cells_frame.pack(fill="x", pady=(10, 0))

        # ---- 底部操作 ----
        bottom = ttk.Frame(main)
        bottom.pack(fill="x", pady=(12, 0))
        self.write_btn = tk.Button(bottom, text="✔ 写入表格", command=lambda: self.do_write(),
                                   font=FONT_BOLD, bg=ACCENT, fg="white",
                                   activebackground="#1d4ed8", activeforeground="white",
                                   relief="flat", bd=0, padx=24, pady=8, cursor="hand2")
        self.write_btn.pack(side="left")
        self.cancel_btn = tk.Button(bottom, text="↺ 取消", command=lambda: self.cancel_edit(),
                                    font=FONT, bg="#e5e7eb", fg=FG,
                                    activebackground="#d1d5db", activeforeground=FG,
                                    relief="flat", bd=0, padx=20, pady=8, cursor="hand2")
        self.cancel_btn.pack(side="left", padx=(8, 0))
        self.ref_btn = tk.Button(bottom, text="◫ 参考", command=self.open_reference,
                                 font=FONT, bg="#e5e7eb", fg=FG,
                                 activebackground="#d1d5db", activeforeground=FG,
                                 relief="flat", bd=0, padx=20, pady=8, cursor="hand2")
        self.ref_btn.pack(side="left", padx=(8, 0))
        self.status_lbl = ttk.Label(bottom, text="", style="Status.TLabel")

        self.switch_combo = ttk.Combobox(bottom, state="readonly", width=9, values=SWITCH_LABELS,
                                         font=FONT_SMALL)
        self.switch_combo.set(SWITCH_LABEL_OF.get(self._switch_key, "询问"))
        self.switch_combo.pack(side="right", padx=(4, 0))
        self.switch_combo.bind("<<ComboboxSelected>>", lambda e: self._save_config())
        ttk.Label(bottom, text="切换行时：").pack(side="right")
        self.status_lbl.pack(side="left", padx=16)

    # ---------------------------------------------------------- 通用

    def _status(self, text: str, kind: str = "normal"):
        style = {"normal": "Status.TLabel", "ok": "Ok.TLabel", "warn": "Warn.TLabel"}[kind]
        self.status_lbl.configure(text="●  " + text, style=style)

    def _autosize(self):
        """按内容多少调整窗口大小（限制在屏幕范围内）。"""
        self.root.update_idletasks()
        w = max(self.root.winfo_reqwidth(), MIN_W)
        h = max(self.root.winfo_reqheight(), MIN_H)
        sw, sh = self.root.winfo_screenwidth(), self.root.winfo_screenheight()
        self.root.geometry(f"{min(w, sw - 60)}x{min(h, sh - 90)}")

    def _selected_index(self) -> int | None:
        sel = self.tree_files.selection()
        if not sel:
            return None
        try:
            return self.tree_files.index(sel[0])
        except Exception:
            return None

    def current_file(self) -> str | None:
        i = self._selected_index()
        if i is not None and 0 <= i < len(self.files):
            return self.files[i]
        return self.files[0] if self.files else None

    @property
    def dirty(self) -> bool:
        """输入框里的内容是否和表格里原本的内容不一样。"""
        return self.values != self.base

    # ---------------------------------------------------------- 配置

    def _window_geometry(self) -> str:
        """当前主窗口的大小和位置（还没显示出来的时候返回空串）。"""
        try:
            g = self.root.geometry()
        except Exception:
            return ""
        return "" if g.startswith("1x1+") else g

    def _save_config(self):
        data = {"files": self.files,
                "last": self.current_file() or "",
                "switch": SWITCH_KEYS.get(self.switch_combo.get(), "ask"),
                "window": self._window_geometry()}
        for p in (CONFIG_FILE, FALLBACK_CONFIG):
            try:
                p.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
                return
            except Exception:
                continue

    def on_close(self):
        """关窗口前把主窗口的位置大小记下来。"""
        self._save_config()
        try:
            self.root.destroy()
        except Exception:
            pass

    # ---------------------------------------------------------- 表格列表

    def _refresh_list(self, keep_index: int | None = None, load: bool = True):
        self.tree_files.delete(*self.tree_files.get_children())
        for f in self.files:
            p = Path(f)
            self.tree_files.insert("", "end", values=("  " + p.name, str(p.parent)))

        if not self.files:
            self.sheet_combo["values"] = []
            self.sheet_var.set("")
            self.values, self.base = {}, {}
            self.loaded_path, self.current_row = None, None
            self.rebuild_cells()
            return

        idx = 0
        if keep_index is not None and 0 <= keep_index < len(self.files):
            idx = keep_index
        elif self.last in self.files:
            idx = self.files.index(self.last)
        items = self.tree_files.get_children()
        if idx < len(items):
            self.tree_files.selection_set(items[idx])
            self.tree_files.focus(items[idx])
            self.tree_files.see(items[idx])
        if load:
            self.on_select_file(force=True)

    def add_file(self):
        path = filedialog.askopenfilename(
            title="选择要填写的表格",
            filetypes=[("Excel 表格", "*.xlsx *.xlsm *.xls *.xlsb *.xltx *.xltm *.xlt *.et"),
                       ("所有文件", "*.*")])
        if not path:
            return
        path = str(Path(path))
        if path not in self.files:
            if not self._can_switch():
                return
            self.files.append(path)
        self.last = path
        self.loaded_path = None
        self._refresh_list(keep_index=self.files.index(path))
        self._save_config()
        self._status(f"已添加：{Path(path).name}", "ok")

    def remove_selected(self):
        i = self._selected_index()
        if i is None or not (0 <= i < len(self.files)):
            self._status("请先选中一条记录", "warn")
            return
        f = self.files[i]
        if not messagebox.askyesno("移出列表", f"把下面这个表格从列表里移除吗？\n\n{f}\n\n（只是移除记录，不会删除文件）"):
            return
        self.files.pop(i)
        if self.loaded_path == f:
            self.loaded_path = None
        self._refresh_list(keep_index=min(i, len(self.files) - 1) if self.files else None)
        self._save_config()
        self._status("已从列表移除", "ok")

    def on_right_click(self, event):
        if not self.files:
            return "break"
        item = self.tree_files.identify_row(event.y)
        if item:
            self.tree_files.selection_set(item)
            self.tree_files.focus(item)
            self.on_select_file(force=True)
        menu = tk.Menu(self.root, tearoff=0, font=FONT)
        menu.add_command(label="打开表格", command=self.open_file)
        menu.add_command(label="打开所在文件夹", command=self.open_folder)
        menu.add_command(label="复制完整路径", command=self.copy_path)
        menu.add_separator()
        menu.add_command(label="从列表中移除记录", command=self.remove_selected)
        try:
            menu.tk_popup(event.x_root, event.y_root)
        finally:
            menu.grab_release()
        return "break"

    def on_select_file(self, event=None, force: bool = False):
        path = self.current_file()
        if not path:
            return
        if path == self.loaded_path and not force:
            return
        if self.dirty and not self._can_switch():
            # 用户放弃了切换，把选中项退回原来的文件
            if self.loaded_path in self.files:
                items = self.tree_files.get_children()
                i = self.files.index(self.loaded_path)
                if i < len(items):
                    self.tree_files.selection_set(items[i])
            return
        self.loaded_path = path
        self.last = path
        self._load_sheets(path)
        self._save_config()

    def on_sheet_change(self):
        if not self._can_switch():
            return
        self.load_current_row()

    def _load_sheets(self, path: str):
        ext = Path(path).suffix.lower()
        names: list[str] = []
        note = ""
        if not Path(path).exists():
            note = "文件不存在，请重新添加"
        elif ext in XLSX_EXTS and openpyxl is not None:
            try:
                wb = openpyxl.load_workbook(path, read_only=True)
                names = list(wb.sheetnames)
                try:
                    biggest = max((wb[n].max_row or 1) for n in names) if names else 1
                    self.row_limit = max(200, min(int(biggest) + 20, 5000))
                except Exception:
                    self.row_limit = 500
                wb.close()
            except Exception as e:
                note = f"读取失败：{e}"
        elif ext in COM_EXTS:
            note = "该格式将在读写时通过 Excel/WPS 处理"
            self.row_limit = 1000

        self.row_combo.configure(values=self._row_options())

        if names:
            self.sheet_combo["values"] = names
            if self.sheet_var.get() not in names:
                self.sheet_var.set(names[0])
        else:
            self.sheet_combo["values"] = []
            self.sheet_var.set("")
        if note:
            self._status(note, "warn")
        self.load_current_row()

    def open_file(self):
        path = self.current_file()
        if not path:
            messagebox.showwarning("提示", "请先添加一个表格。")
            return
        if not Path(path).exists():
            messagebox.showerror("文件不存在", f"找不到文件：\n{path}")
            return
        try:
            os.startfile(path)  # type: ignore[attr-defined]
        except Exception as e:
            messagebox.showerror("打开失败", str(e))

    def open_folder(self):
        path = self.current_file()
        if not path:
            return
        folder = str(Path(path).parent)
        if not Path(folder).exists():
            messagebox.showerror("文件夹不存在", folder)
            return
        try:
            os.startfile(folder)  # type: ignore[attr-defined]
        except Exception as e:
            messagebox.showerror("打开失败", str(e))

    def copy_path(self):
        path = self.current_file()
        if not path:
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(path)
        self._status("路径已复制到剪贴板", "ok")

    # ---------------------------------------------------------- 行切换

    def _row_options(self) -> list[str]:
        """行号下拉里可选的数字。"""
        return [str(i) for i in range(1, self.row_limit + 1)]

    def _reset_row_options(self):
        """每次点开下拉箭头时重建列表，让选择页从头开始显示。"""
        self.row_combo.configure(values=self._row_options())

    def _refresh_ref(self):
        if self.ref_window is not None and self.ref_window.win.winfo_exists():
            self.ref_window.refresh()

    def open_reference(self):
        """「参考」按钮：没开就打开，已经开着就关掉。"""
        if self.ref_window is not None and self.ref_window.win.winfo_exists():
            self.ref_window.close()
            self._status("参考窗口已关闭", "ok")
            return
        self.ref_window = RefWindow(self)
        self._status("参考窗口已打开（位置会记住，直到关掉本程序）")

    def on_row_commit(self, event=None):
        """行号输入框回车 / 选择 / 失去焦点时：确认切换行，并载入该行已有内容。"""
        if self._switching:
            # 询问框一弹出来，输入框就会失焦再触发一次 FocusOut，这里直接挡掉
            return
        try:
            rows = parse_rows(self.row_var.get())
        except ValueError:
            self._status("行号写法无法识别，可用 5 或 5,7 或 5-8", "warn")
            return
        new_row = rows[0] if rows else None
        if new_row == self.current_row:
            return
        if new_row is None:
            return
        if not self._can_switch():
            if self.current_row:
                self.row_var.set(str(self.current_row))
            return
        self.load_current_row()

    def _can_switch(self) -> bool:
        """切换行 / 工作表 / 表格之前调用；返回 False 表示留在原地。"""
        if not self.dirty:
            return True
        if self._switching:
            return True
        # 要写的是“当前显示的这一行”，而不是输入框里刚敲进去的新行号
        target = str(self.current_row) if self.current_row else ""
        if not target:
            return True
        mode = SWITCH_KEYS.get(self.switch_combo.get(), "ask")
        if mode == "write":
            return self.do_write(silent=True, rows_override=target)
        if mode == "discard":
            self.cancel_edit(silent=True)
            return True
        self._switching = True
        try:
            ans = messagebox.askyesnocancel(
                "切换提示",
                f"第 {self.current_row} 行还有没写入的修改。\n\n"
                "是：先写入，再切换\n否：放弃这些修改\n取消：留在当前行",
                parent=self.root)
        finally:
            self._switching = False
        if ans is None:
            return False
        if ans:
            return self.do_write(silent=True, rows_override=target)
        self.cancel_edit(silent=True)
        return True

    def _read_row_data(self, path: str, sheet: str, row: int):
        """读取某一行的现有内容，返回 ({0 基列号: 文本}, 错误信息)。"""
        ext = Path(path).suffix.lower()
        if ext in XLSX_EXTS:
            return read_row_xlsx(path, sheet, row)
        if ext in COM_EXTS:
            if not sheet:
                return {}, "请先选择工作表"
            self.root.config(cursor="watch")
            self.root.update_idletasks()
            try:
                return com_read(path, sheet, row)
            finally:
                self.root.config(cursor="")
        return {}, f"暂不支持 {ext} 格式，请另存为 .xlsx 后再添加"

    def load_current_row(self):
        """读取当前表格 / 工作表 / 行号的内容，显示到输入框里。"""
        try:
            self._load_current_row()
        finally:
            self._refresh_ref()

    def _load_current_row(self):
        path = self.current_file()
        sheet = self.sheet_var.get().strip()
        try:
            rows = parse_rows(self.row_var.get())
        except ValueError:
            rows = []
        row = rows[0] if rows else None
        self.current_row = row

        if not path or row is None or not Path(path).exists():
            self.base, self.values = {}, {}
            self.rebuild_cells()
            return

        data, err = self._read_row_data(path, sheet, row)
        if data is None:
            self.base, self.values = {}, {}
            self.rebuild_cells()
            self._status(err, "warn")
            return

        self.base = dict(data)
        self.values = dict(data)
        # 该行内容比较多时，自动把输入框加够，方便一次看全
        if data:
            need = max(data) + 1 - self.offset
            if need > self.count and need <= MAX_COUNT:
                self.count = need
        self.rebuild_cells()
        if data:
            self._status(f"已载入 第 {row} 行：{len(data)} 列本来就有内容", "ok")
        else:
            self._status(f"第 {row} 行为空，可以直接填写", "ok")

    # ---------------------------------------------------------- 输入框区

    def rebuild_cells(self):
        self._loading = True
        try:
            for w in self.cells_frame.winfo_children():
                w.destroy()
            self.vars = []
            self.entries = []
            self._focus_no_select = None
            if not self.files:
                # 还没有表格：显示一句提示，而不是一片空网格
                ttk.Label(self.cells_frame, text="先在上面「＋ 添加表格」选一个 Excel 文件",
                          style="Empty.TLabel").grid(row=0, column=0, padx=24, pady=34)
                self.range_var.set("")
                self._update_buttons()
                self._autosize()
                return
            for i in range(self.count):
                col0 = self.offset + i
                r, c = divmod(i, PER_ROW)
                holder = ttk.Frame(self.cells_frame, padding=(4, 2))
                holder.grid(row=r, column=c, sticky="nsew", padx=2, pady=2)
                ttk.Label(holder, text=get_column_letter(col0 + 1),
                          style="Col.TLabel", anchor="center").pack(fill="x")
                var = tk.StringVar(value=self.values.get(col0, ""))
                var.trace_add("write", lambda *_a, idx=col0, v=var: self._on_cell_change(idx, v))
                ent = ttk.Entry(holder, textvariable=var, width=10,
                                justify="center", font=FONT)
                ent.pack(fill="x", pady=(2, 0))
                ent.bind("<Return>", lambda e: self.do_write())
                ent.bind("<FocusIn>", self._on_entry_focus)
                ent.bind("<Left>", self._on_left)
                ent.bind("<Right>", self._on_right)
                self.vars.append((col0, var))
                self.entries.append(ent)
        finally:
            self._loading = False
        self._update_range()
        self._update_buttons()
        self._autosize()

    def _on_entry_focus(self, event):
        """输入框获得焦点时，如果里面有内容就全选，方便直接覆盖。"""
        w = event.widget
        if self._focus_no_select is w:
            # 用左右键跳过来的：光标停在文末，不要全选
            self._focus_no_select = None
            return
        w.after_idle(lambda: (w.selection_range(0, "end"), w.icursor("end")))

    # ------------------------------------------------------ 左右键导航

    def _caret(self, w) -> int:
        try:
            return int(w.index("insert"))
        except Exception:
            return 0

    def _focus_cell(self, idx: int, at_end: bool = True) -> bool:
        """把焦点移到第 idx 个输入框，光标停在文末（或文首）。"""
        if not (0 <= idx < len(self.entries)):
            return False
        ent = self.entries[idx]
        if self.root.focus_get() is ent:
            ent.selection_clear()
            ent.icursor("end" if at_end else 0)
            return True
        self._focus_no_select = ent
        ent.focus_set()
        ent.selection_clear()
        ent.icursor("end" if at_end else 0)
        return True

    def _on_left(self, event):
        w = event.widget
        try:
            idx = self.entries.index(w)
        except ValueError:
            return None
        if w.selection_present() or self._caret(w) > 0:
            return None                      # 光标还能往左走，按默认来
        if idx > 0:
            self._focus_cell(idx - 1, at_end=True)
        elif self.offset > 0:
            self.shift(-1, focus_index=idx, at_end=True)
        return "break"

    def _on_right(self, event):
        w = event.widget
        try:
            idx = self.entries.index(w)
        except ValueError:
            return None
        if w.selection_present() or self._caret(w) < len(w.get()):
            return None                      # 光标还能往右走，按默认来
        if idx < len(self.entries) - 1:
            self._focus_cell(idx + 1, at_end=True)
        elif self.offset + self.count < MAX_COL:
            self.shift(1, focus_index=idx, at_end=True)
        return "break"

    def _on_cell_change(self, col0: int, var: tk.StringVar):
        val = var.get()
        if val == "":
            self.values.pop(col0, None)
        else:
            self.values[col0] = val
        self._update_range()

    def _update_range(self):
        first = get_column_letter(self.offset + 1)
        last = get_column_letter(self.offset + self.count)
        hidden = [c for c in set(self.values) | set(self.base)
                  if not (self.offset <= c < self.offset + self.count)]
        extra = f"　（另有 {len(hidden)} 个未显示的列）" if hidden else ""
        self.range_var.set(f"当前可见：{first} 列 → {last} 列{extra}")

    def _update_buttons(self):
        def st(btn, ok):
            btn.configure(state="normal" if ok else "disabled")
        st(self.btn_left, self.offset > 0)
        st(self.btn_right, self.offset + self.count < MAX_COL)
        st(self.btn_add, self.count < MAX_COUNT and self.offset + self.count < MAX_COL)
        st(self.btn_del, self.count > 1)

    def shift(self, delta: int, focus_index: int | None = None, at_end: bool = True):
        new = self.offset + delta
        if new < 0:
            self._status("已经到最左边了（A 列）", "warn")
            return
        if new + self.count > MAX_COL:
            self._status("已经到最右边了", "warn")
            return
        self.offset = new
        self.rebuild_cells()
        if focus_index is not None:
            self._focus_cell(focus_index, at_end=at_end)
        self._status(f"当前最左边的输入框对应 {get_column_letter(self.offset + 1)} 列")

    def reset_offset(self):
        """把输入框重新从第一列（A 列）开始显示。"""
        if self.offset == 0:
            self._status("已经在第一列开头了", "warn")
            return
        self.offset = 0
        self.rebuild_cells()
        self._status("已复位，从 A 列开始显示", "ok")

    def add_cell(self):
        if self.count >= MAX_COUNT:
            self._status(f"最多 {MAX_COUNT} 个输入框，请用「→」移到更高的列", "warn")
            return
        if self.offset + self.count >= MAX_COL:
            self._status("已经到最右边了", "warn")
            return
        self.count += 1
        self.rebuild_cells()

    def del_cell(self):
        if self.count <= 1:
            self._status("至少保留 1 个输入框", "warn")
            return
        removed = self.offset + self.count - 1
        self.count -= 1
        self.rebuild_cells()
        if removed in self.values or removed in self.base:
            self._status(f"{get_column_letter(removed + 1)} 列的内容仍然保留着，点「＋ 增加」可以再看", "warn")

    def cancel_edit(self, silent: bool = False):
        """放弃本次改动，恢复成表格里原来的内容。"""
        self.values = dict(self.base)
        self._loading = True
        try:
            for col0, var in self.vars:
                var.set(self.base.get(col0, ""))
        finally:
            self._loading = False
        self._update_range()
        if not silent:
            self._status("已取消本次改动，恢复为表格里的内容", "ok")

    # ---------------------------------------------------------- 写入

    def do_write(self, silent: bool = False, rows_override: str | None = None) -> bool:
        path = self.current_file()
        if not path:
            if not silent:
                messagebox.showwarning("提示", "请先添加并选择一个表格。")
            return False
        p = Path(path)
        if not p.exists():
            if not silent:
                messagebox.showerror("文件不存在", f"找不到文件：\n{path}\n\n请重新添加。")
            return False
        try:
            rows = parse_rows(rows_override if rows_override is not None else self.row_var.get())
        except ValueError:
            if not silent:
                messagebox.showerror("行号有误", "行号写法无法识别。\n\n可以这样写：5 或 5,7 或 5-8")
            return False
        if not rows:
            if not silent:
                messagebox.showwarning("提示", "请填写要写入的行号。")
            return False

        # 行号框里敲了新行、但还没回车确认：先按“切换行”的规则处理，
        # 免得把上一行的内容写到别的行上去。
        if rows[0] != self.current_row:
            if not self._can_switch():
                return False
            self.load_current_row()
            if not silent:
                self._status(f"已切换到 第 {rows[0]} 行，确认内容后再写入", "warn")
            return False

        # 只写被改动过的列；被清空的列写回空值
        changed = {}
        for col0 in set(self.values) | set(self.base):
            new = self.values.get(col0, "")
            old = self.base.get(col0, "")
            if new != old:
                changed[col0] = new
        if not changed:
            if not silent:
                self._status("没有任何改动，不需要写入", "warn")
            return True

        auto_num = self.auto_num_var.get()
        payload = [(c, (coerce_value(v) if auto_num else (v or None)) if v else None)
                   for c, v in sorted(changed.items())]
        sheet = self.sheet_var.get().strip()
        ext = p.suffix.lower()

        if ext in XLSX_EXTS:
            ok, info = self._write_openpyxl(p, sheet, rows, payload)
        elif ext in COM_EXTS:
            ok, info = self._write_com(p, sheet, rows, payload)
        else:
            ok, info = False, f"暂不支持 {ext} 格式，请另存为 .xlsx 或 .xls 后再添加。"

        if ok:
            self.base = dict(self.values)
            self._status(info, "ok")
        else:
            self._status("写入失败", "warn")
            if not silent:
                messagebox.showerror("写入失败", info)
        return ok

    def _write_openpyxl(self, p: Path, sheet: str, rows: list[int], payload):
        if openpyxl is None:
            if messagebox.askyesno("缺少依赖",
                                   "当前 Python 环境缺少 openpyxl，无法读写 .xlsx。\n\n是否现在自动安装？"):
                self._install_openpyxl()
            return False, "缺少 openpyxl，未写入。"
        if not sheet:
            return False, "请先选择工作表。"
        try:
            wb = openpyxl.load_workbook(str(p), keep_vba=p.suffix.lower() == ".xlsm")
        except PermissionError:
            return False, "无法打开：文件正被其他程序占用，请先关闭 Excel/WPS 里的这个文件。"
        except Exception as e:
            return False, f"打开文件失败：{e}"
        try:
            if sheet not in wb.sheetnames:
                return False, f"工作表「{sheet}」不存在。\n现有工作表：{', '.join(wb.sheetnames)}"
            ws = wb[sheet]
            for r in rows:
                for col0, val in payload:
                    ws.cell(row=r, column=col0 + 1).value = val
            try:
                wb.save(str(p))
            except PermissionError:
                return False, "无法保存：文件正被其他程序占用。\n请先关闭 Excel/WPS 中打开的这个文件，再试一次。"
            except Exception as e:
                return False, f"保存失败：{e}"
        finally:
            try:
                wb.close()
            except Exception:
                pass
        return True, f"已写入 第 {rows_text(rows)} 行，改动 {len(payload)} 个单元格（{p.name}）"

    def _write_com(self, p: Path, sheet: str, rows: list[int], payload):
        if not sheet:
            return False, "请先选择工作表。"
        cells = [{"row": r, "col": col0 + 1, "value": val}
                 for r in rows for col0, val in payload]
        ok, info = com_write(str(p), sheet, cells)
        if ok:
            return True, f"已写入 第 {rows_text(rows)} 行，改动 {len(payload)} 个单元格（{p.name}）"
        return False, info

    def _install_openpyxl(self):
        self._status("正在安装 openpyxl…")
        self.root.update_idletasks()
        flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        try:
            r = subprocess.run([sys.executable, "-m", "pip", "install", "openpyxl"],
                               stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                               text=True, encoding="utf-8", errors="replace",
                               timeout=300, creationflags=flags)
            if r.returncode == 0:
                global openpyxl, get_column_letter
                import openpyxl as _o
                from openpyxl.utils import get_column_letter as _g
                openpyxl, get_column_letter = _o, _g
                messagebox.showinfo("完成", "openpyxl 安装成功，现在可以写入了。")
            else:
                messagebox.showerror("安装失败", (r.stdout or "")[-600:])
        except Exception as e:
            messagebox.showerror("安装失败", f"{e}\n\n可手动执行：{sys.executable} -m pip install openpyxl")


# ---------------------------------------------------------------- 参考浮窗


class RefWindow:
    """只读的“参考”小窗口。

    第一次打开时放在主窗口右边；关掉的时候记住位置，本次运行期间再打开还回到
    上次关闭的地方。重新启动程序后又回到默认位置。
    """

    FLOAT_W = 470
    FLOAT_H = 430

    def __init__(self, app: "FillApp"):
        self.app = app
        self.win = tk.Toplevel(app.root)
        self.win.title("参考（只读）")
        self.win.configure(bg=BG)
        self.win.transient(app.root)

        top = ttk.Frame(self.win, padding=(12, 10, 12, 0))
        top.pack(fill="x")
        ttk.Label(top, text="参考行：").pack(side="left")
        self.row_var = tk.StringVar(value=str(app.current_row or 1))
        self.combo = ttk.Combobox(top, textvariable=self.row_var, width=8, height=12,
                                  justify="center", values=app._row_options())
        self.combo.configure(postcommand=self._reset_options)
        self.combo.pack(side="left")
        self.combo.bind("<Return>", lambda e: self.refresh())
        self.combo.bind("<<ComboboxSelected>>", lambda e: self.refresh())
        self.combo.bind("<FocusOut>", lambda e: self.refresh())
        tk.Button(top, text="✕", command=self.close, font=FONT_BOLD,
                  bg=BG, fg=MUTED, activebackground="#e5e7eb", activeforeground=FG,
                  relief="flat", bd=0, padx=7, cursor="hand2").pack(side="right")
        ttk.Button(top, text="刷新", command=self.refresh, width=5).pack(side="right", padx=(0, 6))

        self.info_var = tk.StringVar(value="")
        self.info_lbl = ttk.Label(self.win, textvariable=self.info_var, style="Hint.TLabel",
                                  anchor="w", wraplength=420)
        self.info_lbl.pack(fill="x", padx=12, pady=(6, 0))

        mid = ttk.Frame(self.win, padding=(12, 8, 12, 12))
        mid.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(mid, columns=("col", "val"), show="headings",
                                 style="Ref.Treeview")
        self.tree.heading("col", text="列")
        self.tree.heading("val", text="内容")
        self.tree.column("col", width=52, anchor="center", stretch=False)
        self.tree.column("val", width=330, anchor="w", stretch=True)
        self.tree.tag_configure("odd", background="#f5f6f8")
        sb = ttk.Scrollbar(mid, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        sb.pack(side="left", fill="y")

        self.win.protocol("WM_DELETE_WINDOW", self.close)
        if app.ref_geom:
            self.win.geometry(self._fit(app.ref_geom))
        else:
            self._default_place()
        self.refresh()

    def _fit(self, geom: str) -> str:
        """把记住的位置拉回屏幕内，免得换了显示器之后跑到看不见的地方。"""
        m = re.fullmatch(r"(\d+)x(\d+)([+-]\d+)([+-]\d+)", geom.strip())
        if not m:
            return geom
        w, h, x, y = (int(m.group(i)) for i in range(1, 5))
        sw, sh = self.win.winfo_screenwidth(), self.win.winfo_screenheight()
        w = min(max(w, 320), max(sw - 20, 320))
        h = min(max(h, 260), max(sh - 60, 260))
        x = min(max(x, 0), max(sw - w - 8, 0))
        y = min(max(y, 0), max(sh - h - 48, 0))
        return f"{w}x{h}+{x}+{y}"

    def _default_place(self):
        """默认位置：贴在主窗口右边；右边放不下就靠屏幕右边。"""
        r = self.app.root
        r.update_idletasks()
        sw, sh = r.winfo_screenwidth(), r.winfo_screenheight()
        w, h = self.FLOAT_W, self.FLOAT_H
        x = r.winfo_rootx() + r.winfo_width() + 12
        if x + w > sw - 10:
            x = max(sw - w - 20, 0)
        y = r.winfo_rooty() + 30
        if y + h > sh - 50:
            y = max(sh - h - 60, 0)
        self.win.geometry(f"{w}x{h}+{x}+{y}")

    def _reset_options(self):
        self.combo.configure(values=self.app._row_options())

    def refresh(self):
        if not self.win.winfo_exists():
            return
        path = self.app.current_file()
        sheet = self.app.sheet_var.get().strip()
        self.tree.delete(*self.tree.get_children())
        if not path:
            self.info_var.set("还没有选择表格。")
            return
        if not Path(path).exists():
            self.info_var.set(f"文件不存在：{path}")
            return
        try:
            rows = parse_rows(self.row_var.get())
        except ValueError:
            rows = []
        row = rows[0] if rows else None
        if row is None:
            self.info_var.set("请填写要参考的行号。")
            return
        data, err = self.app._read_row_data(path, sheet, row)
        if data is None:
            self.info_var.set(err)
            return
        for n, col0 in enumerate(sorted(data)):
            self.tree.insert("", "end", values=(get_column_letter(col0 + 1), data[col0]),
                             tags=("odd",) if n % 2 else ())
        if not data:
            self.tree.insert("", "end", values=("—", "（这一行是空的）"))
        self.info_var.set(f"{Path(path).name} · {sheet or '默认工作表'} · 第 {row} 行"
                          f" · {len(data)} 列有内容")

    def close(self):
        # 记住这次的位置，本次运行期间再打开就回到这里
        try:
            self.app.ref_geom = self.win.geometry()
        except Exception:
            pass
        self.app.ref_window = None
        self.win.destroy()


def _fit_window_geometry(root: tk.Tk, geom: str) -> str:
    """把上次记住的窗口位置拉回屏幕内（换了显示器也不会跑到看不见的地方）。"""
    m = re.fullmatch(r"(\d+)x(\d+)([+-]\d+)([+-]\d+)", (geom or "").strip())
    if not m:
        return ""
    w, h, x, y = (int(m.group(i)) for i in range(1, 5))
    sw, sh = root.winfo_screenwidth(), root.winfo_screenheight()
    w = min(max(w, MIN_W), max(sw - 40, MIN_W))
    h = min(max(h, MIN_H), max(sh - 80, MIN_H))
    x = min(max(x, 0), max(sw - w - 10, 0))
    y = min(max(y, 0), max(sh - h - 50, 0))
    return f"{w}x{h}+{x}+{y}"


def main():
    root = tk.Tk()
    app = FillApp(root)
    root.minsize(MIN_W, MIN_H)
    root.protocol("WM_DELETE_WINDOW", app.on_close)
    root.update_idletasks()

    geom = _fit_window_geometry(root, app.saved_geometry)
    if not geom:
        w = max(root.winfo_reqwidth(), MIN_W)
        h = max(root.winfo_reqheight(), MIN_H)
        x = max((root.winfo_screenwidth() - w) // 2, 0)
        y = max((root.winfo_screenheight() - h) // 3, 0)
        geom = f"{w}x{h}+{x}+{y}"
    root.geometry(geom)
    root.mainloop()


if __name__ == "__main__":
    main()
